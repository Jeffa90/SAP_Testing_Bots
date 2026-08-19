"""Reports: the completed test plan, and the run summary."""

from __future__ import annotations

import pytest
from openpyxl import load_workbook

from saptest.config import load_region
from saptest.core.outcomes import Outcome
from saptest.core.runner import RunConfig, Runner
from saptest.core.status import MessageType, StatusMessage
from saptest.data.workbook import load_dataset
from saptest.drivers.base import VKey
from saptest.drivers.fake import FakeDriver
from saptest.flows import get_flow
from saptest.reporting import write_reports
from saptest.reporting.summary import blocked_by_user, step_failures

PR_NUMBER = "1100712093"


def sap(driver: FakeDriver, key: int) -> None:
    if key == int(VKey.SAVE):
        driver.queue_message(
            StatusMessage(
                MessageType.SUCCESS, "ME", "045",
                f"Purchase requisition number {PR_NUMBER} created", (PR_NUMBER,),
            )
        )


@pytest.fixture
def finished_run(repo_root, tmp_path):
    """A completed two-case run: one clean, one blocked on authorisation."""
    profile = load_region("au", root=repo_root)
    # Point the flow's template at the one real template shipped with the repo.
    profile.templates["standard_po"] = "templates/AU_PTP_Inventory PO.xlsx"

    blocked = {"count": 0}

    def factory():
        driver = FakeDriver(on_press=sap)
        blocked["count"] += 1
        if blocked["count"] == 2:
            driver.queue_message(
                StatusMessage(MessageType.ERROR, text="You are not authorised for ME51N")
            )
        return driver

    config = RunConfig(
        profile=profile,
        flow=get_flow("p2p.pr_create"),
        data=load_dataset(repo_root / "tests" / "fixtures" / "pr_test_data.xlsx", "PR"),
        run_id="reportrun",
        output_root=tmp_path / "runs",
        driver_factory=factory,
        limit=2,
    )
    runner = Runner(config)
    result = runner.run()
    return result, profile, config.flow, runner, tmp_path


# --- the completed test plan --------------------------------------------------------


def test_a_test_plan_is_written_per_case(finished_run, tmp_path):
    result, profile, flow, runner, _ = finished_run
    bundle = write_reports(result, profile, flow, tmp_path / "out", runner.capture)

    assert bundle.problems == []
    assert len(bundle.test_plans) == 2
    assert all(p.exists() for p in bundle.test_plans)


def test_results_are_written_into_the_step_table(finished_run, tmp_path):
    result, profile, flow, runner, _ = finished_run
    bundle = write_reports(result, profile, flow, tmp_path / "out", runner.capture)

    plan = load_workbook(bundle.test_plans[0])["Test Plan"]
    # Template rows 11-13 hold steps 1-3.
    assert plan["G11"].value == "Pass"
    assert PR_NUMBER in str(plan["F11"].value)
    assert plan["G13"].value == "Manual", "the Fiori step must be reported as manual"


def test_an_authorisation_block_is_labelled_as_such_not_as_a_failure(finished_run, tmp_path):
    result, profile, flow, runner, _ = finished_run
    bundle = write_reports(result, profile, flow, tmp_path / "out", runner.capture)

    blocked_case = next(c for c in result.cases if c.outcome is Outcome.BLOCKED_AUTH)
    index = result.cases.index(blocked_case)
    plan = load_workbook(bundle.test_plans[index])["Test Plan"]

    assert plan["G11"].value == "Blocked - No Authorisation"
    assert plan["G12"].value == "Not Executed"


def test_the_header_block_is_filled(finished_run, tmp_path):
    result, profile, flow, runner, _ = finished_run
    bundle = write_reports(
        result, profile, flow, tmp_path / "out", runner.capture,
        header={"business_tester": "J. Judd", "change_control": "CHG0012345"},
    )
    plan = load_workbook(bundle.test_plans[0])["Test Plan"]
    assert plan["G6"].value == "J. Judd"
    assert plan["G2"].value == "CHG0012345"
    assert plan["G1"].value, "test date must be stamped"


def test_screenshots_land_on_the_sheet_named_after_their_step(finished_run, tmp_path):
    result, profile, flow, runner, _ = finished_run
    bundle = write_reports(result, profile, flow, tmp_path / "out", runner.capture)

    workbook = load_workbook(bundle.test_plans[0])
    assert len(workbook["1"]._images) == 2, "step 1 captures 1A and 1B"
    assert len(workbook["2"]._images) == 1, "step 2 captures 2A"


def test_a_rerun_never_overwrites_a_previous_test_plan(finished_run, tmp_path):
    result, profile, flow, runner, _ = finished_run
    first = write_reports(result, profile, flow, tmp_path / "out", runner.capture)
    second = write_reports(result, profile, flow, tmp_path / "out", runner.capture)

    assert set(first.test_plans).isdisjoint(second.test_plans)
    assert all(p.exists() for p in first.test_plans + second.test_plans)


def test_a_missing_template_is_reported_not_raised(finished_run, tmp_path):
    result, profile, flow, runner, _ = finished_run
    profile.templates["standard_po"] = "templates/does-not-exist.xlsx"

    bundle = write_reports(result, profile, flow, tmp_path / "out", runner.capture)

    assert bundle.test_plans == []
    assert any("not found" in p for p in bundle.problems)
    assert bundle.summary_html.exists(), "the summary must still be produced"


# --- the run summary ----------------------------------------------------------------


def test_the_summary_page_is_self_contained(finished_run, tmp_path):
    result, profile, flow, runner, _ = finished_run
    bundle = write_reports(result, profile, flow, tmp_path / "out", runner.capture)

    page = bundle.summary_html.read_text(encoding="utf-8")
    assert "<style>" in page
    assert "http://" not in page and "https://" not in page, "no external requests"
    assert "BLOCKED_AUTH" in page


def test_the_summary_workbook_has_a_sheet_per_view(finished_run, tmp_path):
    result, profile, flow, runner, _ = finished_run
    bundle = write_reports(result, profile, flow, tmp_path / "out", runner.capture)

    workbook = load_workbook(bundle.summary_xlsx)
    assert workbook.sheetnames == ["Cases", "Steps", "Authorisation"]
    assert workbook["Cases"].max_row == 3  # header + two cases


def test_blocked_steps_are_grouped_by_user(finished_run):
    result = finished_run[0]
    grouped = blocked_by_user(result)
    assert grouped
    user, entries = next(iter(grouped.items()))
    assert user
    assert entries[0][1] == "Create Purchase Req."


def test_step_failures_are_counted(finished_run):
    # This run has no FAIL steps -- the blocked one is BLOCKED_AUTH, deliberately.
    assert step_failures(finished_run[0]) == []


def test_summary_json_is_machine_readable(finished_run, tmp_path):
    import json

    result, profile, flow, runner, _ = finished_run
    bundle = write_reports(result, profile, flow, tmp_path / "out", runner.capture)

    data = json.loads(bundle.summary_json.read_text(encoding="utf-8"))
    assert data["flow"] == "p2p.pr_create"
    assert data["cases"][0]["context"]["pr_number"] == PR_NUMBER
    assert data["cases"][0]["steps"][0]["evidence"]


def test_reports_survive_a_run_where_nothing_succeeded(repo_root, tmp_path):
    """An empty or wholly-failed run must still produce a readable report."""
    from saptest.core.outcomes import RunResult

    profile = load_region("au", root=repo_root)
    bundle = write_reports(RunResult("empty", "p2p.pr_create"), profile,
                           get_flow("p2p.pr_create"), tmp_path / "out")
    assert bundle.summary_html.exists()
    assert "No cases were run" in bundle.summary_html.read_text(encoding="utf-8")
