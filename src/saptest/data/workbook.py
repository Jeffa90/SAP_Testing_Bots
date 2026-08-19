"""Reading test data from the Excel workbooks the team already maintains.

Two things this fixes about the original approach:

* **Column names are normalised.** The real workbooks contain ``"Requisition
  Group"`` on one sheet and ``"Requisition  Group"`` (two spaces) on another, and
  the scripts hardcoded whichever spelling their sheet used. Lookup here collapses
  whitespace and ignores case, so both resolve.
* **Values are read as text, exactly as typed.** ``pd.read_excel`` turns a plant
  code into ``1240.0`` and a PR number into a float, which is why the original code
  is littered with ``str(int(row["PO Number"]))``. openpyxl values are converted
  once, here, with integer-valued floats rendered without a decimal tail.
"""

from __future__ import annotations

import logging
import re
from collections.abc import Iterator
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from saptest.core.exceptions import DataError
from saptest.drivers.base import Credentials

log = logging.getLogger(__name__)

_WS = re.compile(r"\s+")


def normalise(column: str) -> str:
    """Canonical form of a column name: collapsed whitespace, lower case."""
    return _WS.sub(" ", str(column or "").strip()).lower()


def cell_text(value: Any, date_format: str = "%d.%m.%Y") -> str:
    """Render a cell as the string SAP should receive.

    Integer-valued floats lose their ``.0``; dates are formatted to the region's
    pattern rather than leaking an ISO timestamp into a SAP date field.
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return "X" if value else ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.strftime(date_format)
    if isinstance(value, date):
        return value.strftime(date_format)
    return str(value).strip()


class Row(dict):
    """One spreadsheet row, addressable by column name however it was spelled."""

    def __init__(self, values: dict[str, str], index: int = 0) -> None:
        super().__init__(values)
        self.index = index
        self._normalised = {normalise(k): k for k in values}

    def get_col(self, column: str, default: str = "") -> str:
        """Look up a column, tolerant of spacing and case."""
        key = self._normalised.get(normalise(column))
        return self.get(key, default) if key else default

    def has(self, column: str) -> bool:
        return normalise(column) in self._normalised

    def first_of(self, *columns: str, default: str = "") -> str:
        """First non-empty value among several candidate column names."""
        for column in columns:
            value = self.get_col(column)
            if value:
                return value
        return default

    def credentials(
        self, which: str = "first", client: str = "", language: str = ""
    ) -> Credentials:
        """Build credentials from the ``First``/``Second`` user columns.

        A blank password is not an error: it means this row logs on via SSO.
        """
        prefix = which.strip().capitalize()
        return Credentials(
            user=self.get_col(f"{prefix} User"),
            password=self.get_col(f"{prefix} Password"),
            client=client,
            language=language,
        )


@dataclass
class DataSet:
    """The rows of one worksheet, plus where they came from."""

    rows: list[Row] = field(default_factory=list)
    columns: list[str] = field(default_factory=list)
    source: Path | None = None
    sheet: str = ""

    def __len__(self) -> int:
        return len(self.rows)

    def __iter__(self) -> Iterator[Row]:
        return iter(self.rows)

    def has_column(self, column: str) -> bool:
        return normalise(column) in {normalise(c) for c in self.columns}

    def require(self, *columns: str) -> None:
        """Fail early and clearly when the sheet is missing something a flow needs."""
        missing = [c for c in columns if not self.has_column(c)]
        if missing:
            raise DataError(
                f"{self.source.name if self.source else 'Workbook'} sheet {self.sheet!r} "
                f"is missing required column(s): {', '.join(missing)}. "
                f"Columns present: {', '.join(self.columns)}"
            )

    def group_by(self, column: str) -> dict[str, list[Row]]:
        """Group rows by a column, preserving sheet order within each group.

        Rows with a blank key are dropped with a warning rather than silently
        forming a phantom group.
        """
        self.require(column)
        groups: dict[str, list[Row]] = {}
        blanks = 0
        for row in self.rows:
            key = row.get_col(column)
            if not key:
                blanks += 1
                continue
            groups.setdefault(key, []).append(row)
        if blanks:
            log.warning("Skipped %d row(s) with a blank %r", blanks, column)
        return groups

    def filter_excluded(self, profile, column: str, kind: str) -> tuple[list[Row], list[Row]]:
        """Split rows into (to run, excluded) using the region's exclusion lists."""
        keep, drop = [], []
        for row in self.rows:
            (drop if profile.is_excluded(kind, row.get_col(column)) else keep).append(row)
        return keep, drop


def load_dataset(
    path: str | Path,
    sheet: str | None = None,
    date_format: str = "%d.%m.%Y",
) -> DataSet:
    """Read one worksheet into a :class:`DataSet`.

    The first row is treated as the header. Rows that are entirely blank are
    dropped, which is what the trailing empty rows in the real workbooks are.
    """
    p = Path(path)
    if not p.exists():
        raise DataError(f"Test data workbook not found: {p}")

    try:
        workbook = load_workbook(p, data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl raises a variety of errors
        raise DataError(f"Could not open workbook {p}: {exc}") from exc

    try:
        if sheet is None:
            worksheet = workbook.worksheets[0]
        elif sheet in workbook.sheetnames:
            worksheet = workbook[sheet]
        else:
            raise DataError(
                f"Sheet {sheet!r} not found in {p.name}. "
                f"Sheets present: {', '.join(workbook.sheetnames)}"
            )

        iterator = worksheet.iter_rows(values_only=True)
        try:
            header = next(iterator)
        except StopIteration:
            raise DataError(f"Sheet {worksheet.title!r} in {p.name} is empty") from None

        columns = [str(h).strip() if h is not None else "" for h in header]
        if not any(columns):
            raise DataError(f"Sheet {worksheet.title!r} in {p.name} has no header row")

        rows: list[Row] = []
        for offset, raw in enumerate(iterator, start=2):
            values = {
                column: cell_text(value, date_format)
                for column, value in zip(columns, raw, strict=False)
                if column
            }
            if any(values.values()):
                rows.append(Row(values, index=offset))

        return DataSet(
            rows=rows, columns=[c for c in columns if c], source=p, sheet=worksheet.title
        )
    finally:
        workbook.close()


def sheet_names(path: str | Path) -> list[str]:
    """Sheet names in a workbook, for the UI's sheet picker."""
    p = Path(path)
    if not p.exists():
        raise DataError(f"Workbook not found: {p}")
    workbook = load_workbook(p, read_only=True)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def load_column_map(*paths: str | Path) -> dict[str, str]:
    """Load the spreadsheet-column -> logical-field map from binding YAML.

    Lives alongside the field bindings because both describe the same mapping
    problem: what a region calls a thing, versus what the flow calls it.
    """
    import yaml  # noqa: PLC0415

    mapping: dict[str, str] = {}
    for path in paths:
        p = Path(path)
        if not p.exists():
            continue
        raw = yaml.safe_load(p.read_text(encoding="utf-8")) or {}
        section = raw.get("columns") if isinstance(raw, dict) else None
        if not section:
            continue
        if not isinstance(section, dict):
            raise DataError(f"'columns' in {p} must be a mapping of column -> field")
        for column, field_name in section.items():
            mapping[normalise(column)] = str(field_name)
    return mapping
