"""Run-level reporting: one HTML page and one workbook covering every case.

The per-case test plans are the formal evidence. This is the operational view: what
passed, what failed, which users were blocked from which steps, and which SAP
messages the catalogue does not yet know about.
"""

from __future__ import annotations

import html
import json
import logging
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from saptest.core.outcomes import Outcome, RunResult

log = logging.getLogger(__name__)

#: Colours shared by the HTML and the workbook, so both read the same way.
_OUTCOME_COLOUR: dict[str, str] = {
    "PASS": "1E7A46",
    "FAIL": "B3261E",
    "ERROR": "8B1A10",
    "BLOCKED_AUTH": "9A6700",
    "MANUAL": "3B5BA5",
    "SKIPPED": "6B7280",
}


def _fmt(value: datetime | None) -> str:
    return value.strftime("%Y-%m-%d %H:%M:%S") if value else ""


def _duration(seconds: float) -> str:
    if seconds < 60:
        return f"{seconds:.0f}s"
    minutes, secs = divmod(int(seconds), 60)
    if minutes < 60:
        return f"{minutes}m {secs}s"
    hours, minutes = divmod(minutes, 60)
    return f"{hours}h {minutes}m"


def blocked_by_user(result: RunResult) -> dict[str, list[tuple[str, str]]]:
    """Authorisation blocks grouped by SAP user, as (case, step name) pairs.

    This is the role-coverage finding the whole exercise exists to produce: which
    users could not perform which steps.
    """
    grouped: dict[str, list[tuple[str, str]]] = defaultdict(list)
    for case in result.cases:
        for step in case.steps:
            if step.outcome is Outcome.BLOCKED_AUTH:
                grouped[case.sap_user or "(unknown user)"].append((case.case_id, step.name))
    return dict(grouped)


def step_failures(result: RunResult) -> list[tuple[str, str, int]]:
    """(step id, step name, failure count), worst first -- where the flow breaks."""
    counts: Counter[tuple[str, str]] = Counter()
    for case in result.cases:
        for step in case.steps:
            if step.outcome in (Outcome.FAIL, Outcome.ERROR):
                counts[(step.step_id, step.name)] += 1
    return [(sid, name, n) for (sid, name), n in counts.most_common()]


# --- workbook ----------------------------------------------------------------------


def write_summary_workbook(result: RunResult, path: Path | str) -> Path:
    """Cross-case summary as a workbook: one row per case, one row per step."""
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="34405A")

    def write_header(sheet, columns: list[str]) -> None:
        sheet.append(columns)
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")
        sheet.freeze_panes = "A2"

    cases = workbook.active
    cases.title = "Cases"
    write_header(cases, ["Case", "SAP User", "Outcome", "Steps Passed", "Duration", "Detail"])
    for case in result.cases:
        passed = sum(1 for s in case.steps if s.outcome is Outcome.PASS)
        cases.append(
            [
                case.case_id,
                case.sap_user,
                case.outcome.value,
                f"{passed}/{len(case.steps)}",
                _duration(case.duration_s),
                case.abort_reason or "; ".join(f"{k}={v}" for k, v in case.context.items()),
            ]
        )
        _colour_outcome(cases.cell(row=cases.max_row, column=3))
    _widths(cases, [16, 16, 16, 14, 12, 70])

    steps = workbook.create_sheet("Steps")
    write_header(
        steps,
        ["Case", "Step", "Name", "Transaction", "Outcome", "Actual Result", "Remediation"],
    )
    for case in result.cases:
        for step in case.steps:
            steps.append(
                [
                    case.case_id,
                    step.step_id,
                    step.name,
                    step.tcode,
                    step.outcome.value,
                    step.actual_result,
                    ", ".join(sorted(set(step.remediations))),
                ]
            )
            _colour_outcome(steps.cell(row=steps.max_row, column=5))
    _widths(steps, [12, 8, 28, 14, 16, 70, 24])

    blocked = workbook.create_sheet("Authorisation")
    write_header(blocked, ["SAP User", "Case", "Blocked Step"])
    for user, entries in sorted(blocked_by_user(result).items()):
        for case_id, step_name in entries:
            blocked.append([user, case_id, step_name])
    _widths(blocked, [20, 14, 40])

    workbook.save(target)
    log.info("Run summary workbook written: %s", target)
    return target


def _colour_outcome(cell) -> None:  # noqa: ANN001
    colour = _OUTCOME_COLOUR.get(str(cell.value), "")
    if colour:
        cell.font = Font(bold=True, color=colour)


def _widths(sheet, widths: list[int]) -> None:  # noqa: ANN001
    from openpyxl.utils import get_column_letter  # noqa: PLC0415

    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


# --- HTML --------------------------------------------------------------------------

_CSS = """
:root{--bg:#fbfbfd;--panel:#fff;--ink:#1a1d23;--muted:#5d6470;--line:#e3e6ec;
--pass:#1e7a46;--fail:#b3261e;--error:#8b1a10;--auth:#9a6700;--manual:#3b5ba5;--skip:#6b7280;}
@media (prefers-color-scheme:dark){:root{--bg:#14161a;--panel:#1c1f25;--ink:#e8eaee;
--muted:#9aa1ad;--line:#2c313a;--pass:#57c98a;--fail:#f2837a;--error:#ff9d92;
--auth:#e3b341;--manual:#8fb0f0;--skip:#8b929e;}}
*{box-sizing:border-box}
body{margin:0;padding:2rem 1.5rem;background:var(--bg);color:var(--ink);
font:15px/1.55 -apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,sans-serif;}
.wrap{max-width:1100px;margin:0 auto}
h1{font-size:1.5rem;margin:0 0 .25rem}
.sub{color:var(--muted);margin:0 0 1.75rem;font-size:.92rem}
.cards{display:flex;flex-wrap:wrap;gap:.75rem;margin-bottom:1.75rem}
.card{background:var(--panel);border:1px solid var(--line);border-radius:10px;
padding:.8rem 1.1rem;min-width:118px}
.card .n{font-size:1.7rem;font-weight:650;line-height:1.1}
.card .l{font-size:.72rem;letter-spacing:.07em;text-transform:uppercase;color:var(--muted)}
section{background:var(--panel);border:1px solid var(--line);border-radius:10px;
padding:1.1rem 1.25rem;margin-bottom:1.25rem}
h2{font-size:1rem;margin:0 0 .85rem;letter-spacing:.01em}
.scroll{overflow-x:auto}
table{border-collapse:collapse;width:100%;font-size:.88rem;min-width:640px}
th{text-align:left;font-weight:600;color:var(--muted);font-size:.74rem;
letter-spacing:.06em;text-transform:uppercase;padding:.4rem .65rem;
border-bottom:1px solid var(--line)}
td{padding:.48rem .65rem;border-bottom:1px solid var(--line);vertical-align:top}
tr:last-child td{border-bottom:none}
.PASS{color:var(--pass);font-weight:600}.FAIL{color:var(--fail);font-weight:600}
.ERROR{color:var(--error);font-weight:600}.BLOCKED_AUTH{color:var(--auth);font-weight:600}
.MANUAL{color:var(--manual);font-weight:600}.SKIPPED{color:var(--skip)}
code{font:.85em ui-monospace,SFMono-Regular,Menlo,monospace;background:var(--bg);
padding:.1em .35em;border-radius:4px}
.note{color:var(--muted);font-size:.86rem;margin:.4rem 0 0}
.empty{color:var(--muted);font-style:italic}
"""


def write_run_summary(
    result: RunResult, path: Path | str, unknown: list[dict] | None = None
) -> Path:
    """Write the run summary as a self-contained HTML page."""
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(_render(result, unknown or []), encoding="utf-8")
    log.info("Run summary written: %s", target)
    return target


def _esc(value: object) -> str:
    return html.escape(str(value if value is not None else ""))


def _aborted_block(result: RunResult) -> str:
    if not result.aborted:
        return ""
    return f"<section><h2>Run stopped early</h2><p>{_esc(result.abort_reason)}</p></section>"


def _render(result: RunResult, unknown: list[dict]) -> str:
    tally = result.tally()
    cards = "".join(
        f'<div class="card"><div class="n {o.value}">{n}</div>'
        f'<div class="l">{o.value.replace("_", " ").title()}</div></div>'
        for o, n in tally.items()
        if n or o in (Outcome.PASS, Outcome.FAIL)
    )

    case_rows = "".join(
        "<tr>"
        f"<td><code>{_esc(c.case_id)}</code></td>"
        f"<td>{_esc(c.sap_user)}</td>"
        f'<td class="{c.outcome.value}">{c.outcome.value}</td>'
        f"<td>{sum(1 for s in c.steps if s.outcome is Outcome.PASS)}/{len(c.steps)}</td>"
        f"<td>{_duration(c.duration_s)}</td>"
        f"<td>{_esc(c.abort_reason or c.context.get('pr_number') or '')}</td>"
        "</tr>"
        for c in result.cases
    )

    failures = step_failures(result)
    failure_rows = "".join(
        f"<tr><td>{_esc(sid)}</td><td>{_esc(name)}</td><td>{n}</td></tr>"
        for sid, name, n in failures
    )

    blocked = blocked_by_user(result)
    blocked_rows = "".join(
        f"<tr><td>{_esc(user)}</td><td>{len(entries)}</td>"
        f"<td>{_esc(', '.join(sorted({s for _, s in entries})))}</td></tr>"
        for user, entries in sorted(blocked.items())
    )

    unknown_rows = "".join(
        f"<tr><td><code>{_esc(u.get('identity') or u.get('key'))}</code></td>"
        f"<td>{_esc(u.get('count', 1))}</td>"
        f"<td>{_esc(', '.join(u.get('tcodes', [])))}</td>"
        f"<td>{_esc(u.get('sample_text', ''))}</td></tr>"
        for u in unknown
    )

    def table(headers: list[str], rows: str, empty: str) -> str:
        if not rows:
            return f'<p class="empty">{_esc(empty)}</p>'
        head = "".join(f"<th>{h}</th>" for h in headers)
        return (
            '<div class="scroll"><table><thead><tr>'
            f"{head}</tr></thead><tbody>{rows}</tbody></table></div>"
        )

    return f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>{_esc(result.run_id)} - SAP test run</title>
<style>{_CSS}</style></head><body><div class="wrap">

<h1>{_esc(result.flow)} &middot;
<span class="{result.outcome.value}">{result.outcome.value}</span></h1>
<p class="sub">Run <code>{_esc(result.run_id)}</code> &middot; region {_esc(result.region)}
&middot; system {_esc(result.system)} &middot; started {_esc(_fmt(result.started_at))}
&middot; {_duration(result.duration_s)}</p>

<div class="cards">{cards}</div>

{_aborted_block(result)}

<section><h2>Cases</h2>
{table(["Case", "SAP user", "Outcome", "Steps passed", "Duration", "Detail"],
       case_rows, "No cases were run.")}
</section>

<section><h2>Authorisation blocks by user</h2>
{table(["SAP user", "Blocked steps", "Which"], blocked_rows,
       "No authorisation blocks: every user could perform every step attempted.")}
<p class="note">A blocked step means the user lacked the authorisation to perform it.
That is a role finding, not a defect &mdash; the step was never tested.</p>
</section>

<section><h2>Where the flow failed</h2>
{table(["Step", "Name", "Failures"], failure_rows, "No steps failed.")}
</section>

<section><h2>Messages not yet in the catalogue</h2>
{table(["Message", "Seen", "Transactions", "Text"], unknown_rows,
       "Every SAP message encountered was recognised by the error catalogue.")}
<p class="note">Run <code>saptest catalog review</code> to turn these into catalogue
entries, so the next run knows how to handle them.</p>
</section>

</div></body></html>"""


def summary_json(result: RunResult) -> str:
    """Machine-readable summary, for the UI's run history."""
    return json.dumps(
        {
            "run_id": result.run_id,
            "flow": result.flow,
            "region": result.region,
            "system": result.system,
            "outcome": result.outcome.value,
            "started_at": _fmt(result.started_at),
            "duration_s": round(result.duration_s, 1),
            "aborted": result.aborted,
            "tally": {k.value: v for k, v in result.tally().items()},
            "cases": [
                {
                    "case_id": c.case_id,
                    "sap_user": c.sap_user,
                    "outcome": c.outcome.value,
                    "context": c.context,
                    "steps": [
                        {
                            "step_id": s.step_id,
                            "name": s.name,
                            "tcode": s.tcode,
                            "outcome": s.outcome.value,
                            "actual_result": s.actual_result,
                            "evidence": [e.path for e in s.evidence],
                        }
                        for s in c.steps
                    ],
                }
                for c in result.cases
            ],
        },
        indent=2,
    )
